from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook

from excel_dbapi.engines.openpyxl.backend import OpenpyxlBackend
from excel_dbapi.engines.result import ExecutionResult
from excel_dbapi.exceptions import ProgrammingError
from excel_dbapi.executor import SharedExecutor
from excel_dbapi.parser import (
    _OrderByClause,
    _bind_expression_values,
    _bind_params,
    _collect_case_tokens_until,
    _expression_to_sql_for_order_by,
    _expression_values_to_bind,
    _parse_case_expression_tokens,
    _parse_column_expression,
    _parse_compound,
    _parse_order_by_clause_tokens,
    _parse_order_by_item_tokens,
    _tokenize_expression,
    _tokenize,
    _validate_join_column_reference,
    _validate_join_where_node,
    _where_to_sql_for_order_by,
    parse_sql,
)


def _make_join_workbook(path: Path) -> OpenpyxlBackend:
    workbook = Workbook()
    left = workbook.active
    assert left is not None
    left.title = "t1"
    left.append(["id", "name", "value"])
    left.append([1, "A", 10])
    left.append([2, "B", 20])

    right = workbook.create_sheet("t2")
    right.append(["id", "label", "score"])
    right.append([1, "x", 5])
    right.append([2, "y", 7])

    workbook.save(path)
    return OpenpyxlBackend(str(path))


def test_parser_order_by_clause_string_index_guards_multiple_items() -> None:
    clause = _OrderByClause([
        {"column": "a", "direction": "ASC"},
        {"column": "b", "direction": "DESC"},
    ])
    with pytest.raises(TypeError, match="list indices must be integers"):
        _ = clause["column"]


def test_parser_order_by_clause_string_index_single_item() -> None:
    clause = _OrderByClause([{"column": "a", "direction": "ASC"}])
    assert clause["column"] == "a"


def test_parser_invalid_join_keyword_variants() -> None:
    with pytest.raises(ValueError, match="Unsupported SQL syntax: RIGHT"):
        parse_sql("SELECT t1.id FROM t1 RIGHT t2 ON t1.id = t2.id")
    with pytest.raises(ValueError, match="Unsupported SQL syntax: FULL"):
        parse_sql("SELECT t1.id FROM t1 FULL t2 ON t1.id = t2.id")
    with pytest.raises(ValueError, match="Unsupported SQL syntax: CROSS"):
        parse_sql("SELECT t1.id FROM t1 CROSS t2")


def test_parser_invalid_alter_table_variants() -> None:
    with pytest.raises(ValueError, match="Invalid ALTER TABLE format"):
        parse_sql("ALTER TABLE t1")
    with pytest.raises(ValueError, match="Invalid ALTER TABLE format"):
        parse_sql("ALTER TABLE t1 DROP x")
    with pytest.raises(ValueError, match="Invalid ALTER TABLE format"):
        parse_sql("ALTER TABLE t1 RENAME COLUMN a b")


def test_parser_compound_offset_limit_placeholders_missing_params() -> None:
    with pytest.raises(ValueError, match="Not enough parameters for LIMIT placeholder"):
        parse_sql("SELECT id FROM t1 UNION SELECT id FROM t2 LIMIT ?", ())
    with pytest.raises(ValueError, match="Not enough parameters for OFFSET placeholder"):
        parse_sql("SELECT id FROM t1 UNION SELECT id FROM t2 OFFSET ?", ())


def test_parse_compound_rejects_non_select_branch() -> None:
    with pytest.raises(ValueError, match="Compound queries support only SELECT subqueries"):
        _parse_compound("SELECT id FROM t1 UNION (UPDATE t2 SET id = 1)", None)


def test_execute_compound_structure_guards(tmp_path: Path) -> None:
    backend = _make_join_workbook(tmp_path / "compound_guards.xlsx")
    executor = SharedExecutor(backend)
    with pytest.raises(ValueError, match="at least one SELECT"):
        executor._execute_compound({"queries": []})
    with pytest.raises(ValueError, match="valid operator"):
        executor._execute_compound({"queries": [{"action": "SELECT"}]})
    with pytest.raises(ValueError, match="Invalid COMPOUND query structure"):
        executor._execute_compound(
            {
                "queries": [
                    parse_sql("SELECT id FROM t1"),
                    parse_sql("SELECT id FROM t2"),
                ],
                "operators": [],
            }
        )
    backend.close()


def test_execute_compound_invalid_operator_and_order_resolution(tmp_path: Path) -> None:
    backend = _make_join_workbook(tmp_path / "compound_ops.xlsx")
    executor = SharedExecutor(backend)

    with pytest.raises(ValueError, match="Unsupported compound operator"):
        executor.execute(
            {
                "action": "COMPOUND",
                "operators": ["MERGE"],
                "queries": [parse_sql("SELECT id FROM t1"), parse_sql("SELECT id FROM t2")],
            }
        )

    with pytest.raises(ValueError, match="ORDER BY column 'missing' not found"):
        executor.execute(
            parse_sql("SELECT id FROM t1 UNION SELECT id FROM t2 ORDER BY missing")
        )

    rows = executor.execute(parse_sql("SELECT id FROM t1 UNION ALL SELECT id FROM t2 OFFSET 1")).rows
    assert rows == [(2,), (1,), (2,)]
    backend.close()


def test_executor_expression_serialization_case_and_where_shapes() -> None:
    simple_case = {
        "type": "case",
        "mode": "simple",
        "value": {"type": "column", "source": "t1", "name": "id"},
        "whens": [{"match": {"type": "literal", "value": 1}, "result": {"type": "literal", "value": "one"}}],
        "else": {"type": "literal", "value": "other"},
    }
    searched_case = {
        "type": "case",
        "mode": "searched",
        "whens": [
            {
                "condition": {
                    "type": "not",
                    "operand": {
                        "type": "compound",
                        "conditions": [{"column": "t1.id", "operator": "IN", "value": [1, 2]}],
                        "conjunctions": ["AND"],
                    },
                },
                "result": {"type": "literal", "value": "x"},
            }
        ],
        "else": {"type": "literal", "value": "y"},
    }

    assert SharedExecutor._output_name(simple_case) == "case_expr"
    assert SharedExecutor._source_key({"type": "alias", "alias": "k", "expression": simple_case}).startswith("__expr__:")
    assert SharedExecutor._expression_to_sql({"type": "alias", "expression": {"type": "aggregate", "func": "COUNT", "arg": "*"}}) == "COUNT(*)"

    assert "NOT" in SharedExecutor._where_to_sql({"type": "not", "operand": {"column": "id", "operator": "=", "value": 1}})
    assert SharedExecutor._where_to_sql({"type": "not", "operand": "bad"}) == "NOT"
    assert SharedExecutor._where_to_sql({"type": "compound", "conditions": [{"column": "id", "operator": "=", "value": 1}], "conjunctions": ["AND"]}).startswith("(")
    assert "(SUBQUERY)" in SharedExecutor._where_operand_to_sql({"type": "subquery"}, is_column=False)
    assert "BETWEEN" in SharedExecutor._where_to_sql({"column": "id", "operator": "BETWEEN", "value": [1]})
    assert "CASE" in SharedExecutor._expression_to_sql(searched_case)


def test_executor_collect_expression_and_where_refs_paths() -> None:
    expr = {
        "type": "alias",
        "alias": "z",
        "expression": {
            "type": "binary_op",
            "op": "+",
            "left": {"type": "column", "source": "t1", "name": "id"},
            "right": {"type": "literal", "value": 1},
        },
    }
    where = {
        "type": "not",
        "operand": {
            "conditions": [
                {
                    "column": {"type": "column", "source": "t1", "name": "id"},
                    "operator": "IN",
                    "value": [
                        {"type": "column", "source": "t2", "name": "id"},
                        {"type": "literal", "value": 2},
                    ],
                }
            ]
        },
    }

    assert SharedExecutor._contains_arithmetic_expression({"type": "alias", "alias": "a", "expression": {"type": "alias", "alias": "b", "expression": expr}})
    assert SharedExecutor._collect_expression_column_refs("t1.id") == {"t1.id"}
    assert SharedExecutor._collect_expression_column_refs(3.14) == set()
    refs = SharedExecutor._collect_where_column_refs(where)
    assert "t1.id" in refs
    assert "t2.id" in refs


def test_executor_apply_order_by_and_sort_key_edges(tmp_path: Path) -> None:
    backend = _make_join_workbook(tmp_path / "sort_edges.xlsx")
    executor = SharedExecutor(backend)

    assert executor._apply_order_by([], None, value_getter=lambda row, col: None) == []

    single = [{"id": 1}]
    ordered = executor._apply_order_by(
        single,
        [{"column": "id", "direction": "ASC"}],
        value_getter=lambda row, col: row[col],
        available_columns={"id"},
    )
    assert ordered == single

    assert executor._sort_key(None) == (1, (0, ""))
    assert executor._sort_key("12") == (0, (0, 12.0))
    assert executor._sort_key(True) == (0, (1, 1))
    backend.close()


def test_executor_validate_join_where_node_case_branches() -> None:
    seen: list[tuple[str, str, str]] = []

    def _validate(source: str, name: str, context: str) -> None:
        seen.append((source, name, context))

    node = {
        "type": "not",
        "operand": {
            "conditions": [
                {
                    "column": {
                        "type": "case",
                        "mode": "searched",
                        "whens": [
                            "skip",
                            {
                                "condition": {
                                    "column": {"type": "column", "source": "t1", "name": "id"},
                                    "operator": "=",
                                    "value": {"type": "column", "source": "t2", "name": "id"},
                                },
                                "result": {"type": "literal", "value": 1},
                            },
                        ],
                    },
                    "operator": "IN",
                    "value": [
                        {"type": "alias", "alias": "a", "expression": {"type": "column", "source": "t2", "name": "score"}},
                        None,
                    ],
                }
            ]
        },
    }
    SharedExecutor._validate_join_where_node(node, _validate)
    assert ("t1", "id", "WHERE") in seen
    assert ("t2", "id", "WHERE") in seen
    assert ("t2", "score", "WHERE") in seen


def test_executor_join_validation_and_case_expression_errors(tmp_path: Path) -> None:
    backend = _make_join_workbook(tmp_path / "join_validate.xlsx")
    executor = SharedExecutor(backend)

    with pytest.raises(ValueError, match="Aggregate arguments in JOIN queries must be qualified"):
        executor.execute(parse_sql("SELECT SUM(id) FROM t1 a JOIN t2 b ON a.id = b.id"))

    with pytest.raises(ValueError, match="requires qualified column names in JOIN queries"):
        executor.execute(parse_sql("SELECT CASE id WHEN 1 THEN 'x' ELSE 'y' END FROM t1 a JOIN t2 b ON a.id = b.id"))

    with pytest.raises(ValueError, match="Invalid source reference in WHERE"):
        executor.execute(parse_sql("SELECT a.id FROM t1 a JOIN t2 b ON a.id = b.id WHERE c.id = 1"))

    backend.close()


def test_executor_eval_expression_and_condition_error_paths(tmp_path: Path) -> None:
    backend = _make_join_workbook(tmp_path / "eval_expr.xlsx")
    executor = SharedExecutor(backend)
    row = {"t1.id": 1, "t1.value": "text", "x": None}

    with pytest.raises(ProgrammingError, match="Unsupported unary operator"):
        executor._eval_expression({"type": "unary_op", "op": "~", "operand": {"type": "literal", "value": 1}}, row, lambda c: row.get(c))
    with pytest.raises(ProgrammingError, match="requires numeric operands"):
        executor._eval_expression({"type": "binary_op", "op": "+", "left": {"type": "literal", "value": "a"}, "right": {"type": "literal", "value": "b"}}, row, lambda c: row.get(c))
    with pytest.raises(ProgrammingError, match="Unsupported arithmetic operator"):
        executor._eval_expression({"type": "binary_op", "op": "%", "left": {"type": "literal", "value": 3}, "right": {"type": "literal", "value": 2}}, row, lambda c: row.get(c))
    with pytest.raises(ProgrammingError, match="Aggregate expressions are not supported"):
        executor._eval_expression({"type": "aggregate", "func": "COUNT", "arg": "*"}, row, lambda c: row.get(c))
    with pytest.raises(ProgrammingError, match="Unsupported expression type"):
        executor._eval_expression({"type": "mystery"}, row, lambda c: row.get(c))

    assert executor._evaluate_condition(row, {"column": "x", "operator": "NOT IN", "value": (1, 2)}) is False
    assert executor._evaluate_condition(row, {"column": "x", "operator": "NOT BETWEEN", "value": (1, 2)}) is False
    assert executor._evaluate_condition({"x": 5}, {"column": "x", "operator": "NOT BETWEEN", "value": (None, 2)}) is False
    with pytest.raises(NotImplementedError, match="Unsupported LIKE pattern type"):
        executor._evaluate_condition({"x": "abc"}, {"column": "x", "operator": "NOT LIKE", "value": 5})

    backend.close()


def test_executor_aggregate_edge_paths(tmp_path: Path) -> None:
    backend = _make_join_workbook(tmp_path / "agg_edges.xlsx")
    executor = SharedExecutor(backend)

    with pytest.raises(ValueError, match=r"COUNT\(DISTINCT \*\) is not supported"):
        executor._compute_aggregate("COUNT", "*", [{"*": 1}], distinct=True)
    with pytest.raises(ValueError, match="Unsupported aggregate function"):
        executor._compute_aggregate("MEDIAN", "id", [{"id": 1}], distinct=False)
    assert executor._aggregate_spec_from_label("COUNT()") is None
    assert executor._aggregate_spec_from_label("SUM(DISTINCT t1.id)") is None

    with pytest.raises(ValueError, match="must appear in GROUP BY"):
        executor.execute(parse_sql("SELECT id, COUNT(*) FROM t1 GROUP BY name"))
    with pytest.raises(ValueError, match="must be a GROUP BY column"):
        executor.execute(parse_sql("SELECT name, COUNT(*) FROM t1 GROUP BY name HAVING id > 1"))

    backend.close()


def test_executor_compound_manual_description_mismatch(tmp_path: Path) -> None:
    backend = _make_join_workbook(tmp_path / "compound_desc_mismatch.xlsx")

    class _FakeExecutor(SharedExecutor):
        def __init__(self, backend_obj: OpenpyxlBackend) -> None:
            super().__init__(backend_obj)
            self._idx = 0

        def execute(self, parsed: dict[str, Any], **kwargs: Any) -> ExecutionResult:  # type: ignore[override]
            self._idx += 1
            if self._idx == 1:
                return ExecutionResult("SELECT", [(1,)], [("id", None, None, None, None, None, None)], 1)
            return ExecutionResult(
                "SELECT",
                [(1, "x")],
                [("id", None, None, None, None, None, None), ("name", None, None, None, None, None, None)],
                1,
            )

    fake = _FakeExecutor(backend)
    with pytest.raises(ValueError, match="matching column counts"):
        fake._execute_compound(
            {
                "queries": [{"action": "SELECT"}, {"action": "SELECT"}],
                "operators": ["UNION"],
            }
        )
    backend.close()


def test_parser_column_expression_internal_error_paths() -> None:
    with pytest.raises(ValueError, match="Invalid column expression"):
        _parse_column_expression("   ")
    with pytest.raises(ValueError, match="aggregate functions are not supported"):
        _parse_column_expression("COUNT(id)", allow_aggregates=False)
    with pytest.raises(ValueError, match="Unsupported function: COUNT"):
        _parse_column_expression("COUNT()")
    with pytest.raises(ValueError, match="wildcard is not supported"):
        _parse_column_expression("*", allow_wildcard=False)
    with pytest.raises(ValueError, match="Unsupported column expression"):
        _parse_column_expression("(a + 1")
    with pytest.raises(ValueError, match="cannot be used inside arithmetic expressions"):
        _parse_column_expression("COUNT(id) + 1")


def test_parser_case_expression_token_errors() -> None:
    with pytest.raises(ValueError, match="Invalid CASE expression"):
        _parse_case_expression_tokens(_tokenize("x"), 0)
    with pytest.raises(ValueError, match="missing WHEN"):
        _parse_case_expression_tokens(_tokenize("CASE ELSE 1 END"), 0)
    with pytest.raises(ValueError, match="THEN requires a result"):
        _parse_case_expression_tokens(_tokenize("CASE WHEN a = 1 THEN END"), 0)
    with pytest.raises(ValueError, match="ELSE requires a result"):
        _parse_case_expression_tokens(_tokenize("CASE WHEN a = 1 THEN 1 ELSE END"), 0)
    with pytest.raises(ValueError, match="missing END"):
        _parse_case_expression_tokens(_tokenize("CASE WHEN a = 1 THEN 1"), 0)


def test_parser_order_by_internal_sql_rendering_branches() -> None:
    where_compound = {
        "type": "compound",
        "conditions": [
            "a.id = 1",
            {"column": "a.id", "operator": "NOT IN", "value": [{"type": "literal", "value": "x"}]},
        ],
        "conjunctions": ["OR"],
    }
    where_not = {"type": "not", "operand": {"column": "a.id", "operator": "BETWEEN", "value": [1, 2]}}
    where_not_bad = {"type": "not", "operand": "oops"}
    assert _where_to_sql_for_order_by(where_compound).startswith("(")
    assert _where_to_sql_for_order_by(where_not).startswith("NOT")
    assert _where_to_sql_for_order_by(where_not_bad) == "NOT"
    assert "BETWEEN" in _where_to_sql_for_order_by({"column": "a.id", "operator": "BETWEEN", "value": [1]})
    assert "IN" in _where_to_sql_for_order_by({"column": "a.id", "operator": "IN", "value": 3})

    expr = {
        "type": "case",
        "mode": "searched",
        "whens": [
            "skip",
            {
                "condition": {"column": "a.id", "operator": "=", "value": 1},
                "result": {"type": "literal", "value": "one"},
            },
        ],
        "else": {"type": "literal", "value": "other"},
    }
    assert "CASE" in _expression_to_sql_for_order_by(expr)


def test_parser_order_by_item_and_clause_validation_paths() -> None:
    with pytest.raises(ValueError, match="Invalid ORDER BY direction"):
        _parse_order_by_item_tokens(_tokenize("CASE WHEN a = 1 THEN 1 END SIDEWAYS"))
    with pytest.raises(ValueError, match="Unsupported SQL syntax"):
        _parse_order_by_item_tokens(_tokenize("CASE WHEN a = 1 THEN 1 END ASC EXTRA"))
    with pytest.raises(ValueError, match="Unsupported SQL syntax"):
        _parse_order_by_item_tokens(["a", "ASC", "EXTRA"])
    with pytest.raises(ValueError, match="Invalid ORDER BY clause format"):
        _parse_order_by_clause_tokens([])
    parsed = _parse_order_by_clause_tokens(_tokenize("a DESC, b ASC"))
    assert parsed[0]["column"] == "a"
    assert parsed[1]["column"] == "b"


def test_parser_select_clause_format_errors() -> None:
    with pytest.raises(ValueError, match="Invalid SQL query format"):
        parse_sql("SELECT FROM t1")
    with pytest.raises(ValueError, match="Expected alias after AS"):
        parse_sql("SELECT id FROM t1 AS")


def test_executor_join_helper_edges_without_upsert(tmp_path: Path) -> None:
    backend = _make_join_workbook(tmp_path / "join_helpers.xlsx")
    executor = SharedExecutor(backend)

    seen: list[tuple[str, str, str]] = []

    def _validate(source: str, name: str, context: str) -> None:
        seen.append((source, name, context))

    SharedExecutor._validate_join_where_refs(
        {
            "conditions": [
                {
                    "column": {"type": "unary_op", "op": "-", "operand": {"type": "column", "source": "a", "name": "id"}},
                    "operator": "=",
                    "value": {
                        "type": "binary_op",
                        "op": "+",
                        "left": {"type": "column", "source": "b", "name": "id"},
                        "right": {"type": "literal", "value": 1},
                    },
                }
            ]
        },
        _validate,
    )
    assert ("a", "id", "WHERE") in seen
    assert ("b", "id", "WHERE") in seen

    with pytest.raises(ValueError, match="Unknown source reference"):
        executor._resolve_join_column({}, {"source": "missing", "name": "id"})
    flattened = executor._flatten_join_row({"a": {"id": 1}, "bad": 42})
    assert flattened == {"a.id": 1}
    backend.close()


def test_parser_tokenize_expression_and_case_collection_edges() -> None:
    tokens = _tokenize_expression("'it''s' + \"a\"\"b\" + (x)")
    assert "'it''s'" in tokens
    assert '\"a\"\"b\"' in tokens

    collected, index, stop = _collect_case_tokens_until(
        _tokenize("CASE WHEN (a = 1) THEN 'x' END"),
        1,
        {"THEN"},
    )
    assert stop == "THEN"
    assert index > 1
    assert collected


def test_parser_expression_binding_case_paths() -> None:
    expression = {
        "type": "alias",
        "alias": "expr",
        "expression": {
            "type": "case",
            "mode": "simple",
            "value": {"type": "literal", "value": "?"},
            "whens": [
                "skip",
                {
                    "match": {"type": "literal", "value": "?"},
                    "result": {
                        "type": "binary_op",
                        "op": "+",
                        "left": {"type": "literal", "value": "?"},
                        "right": {"type": "literal", "value": 1},
                    },
                },
            ],
            "else": {"type": "unary_op", "op": "-", "operand": {"type": "literal", "value": "?"}},
        },
    }
    values = _expression_values_to_bind(expression)
    assert values == ["?", "?", "?", 1, "?"]
    consumed = _bind_expression_values(expression, [10, 20, 30, 40, 50], 0)
    assert consumed == 5


def test_parser_bind_params_too_many_values() -> None:
    with pytest.raises(ValueError, match="Too many parameters for placeholders"):
        _bind_params([1], (9,))


def test_parser_insert_multi_row_scanner_edges() -> None:
    parsed = parse_sql("INSERT INTO t VALUES ((1)), (2)")
    assert parsed["values"] == [["(1)"], [2]]

    with pytest.raises(ValueError, match="Invalid INSERT format"):
        parse_sql("INSERT INTO t VALUES (1, 'x)")

    with pytest.raises(ValueError, match="Too many parameters for placeholders"):
        parse_sql("INSERT INTO t VALUES (?, ?), (?, ?)", (1, 2, 3, 4, 5))


def test_parser_alter_valid_paths_and_compound_parenthesized_branches() -> None:
    assert parse_sql("ALTER TABLE t ADD COLUMN c FLOAT")["type_name"] == "REAL"
    assert parse_sql("ALTER TABLE t DROP COLUMN c")["operation"] == "DROP_COLUMN"
    assert parse_sql("ALTER TABLE t RENAME COLUMN c TO d")["operation"] == "RENAME_COLUMN"

    compound = parse_sql("(SELECT id FROM t1) UNION (SELECT id FROM t2 ORDER BY id DESC LIMIT 2 OFFSET 1)")
    assert compound["action"] == "COMPOUND"


def test_parser_compound_trailing_clause_validation_errors() -> None:
    with pytest.raises(ValueError, match="Invalid LIMIT clause format"):
        parse_sql("SELECT id FROM t1 UNION SELECT id FROM t2 LIMIT")
    with pytest.raises(ValueError, match="Invalid OFFSET clause format"):
        parse_sql("SELECT id FROM t1 UNION SELECT id FROM t2 OFFSET")

    parsed = parse_sql("SELECT id FROM t1 UNION SELECT id FROM t2 LIMIT ? OFFSET ?", (2, 1))
    assert parsed["limit"] == 2
    assert parsed["offset"] == 1


def test_parser_compound_internal_malformed_shapes() -> None:
    assert _parse_compound("   ", None) is None
    assert _parse_compound("UPDATE t SET x = 1", None) is None

    with pytest.raises(ValueError, match="Invalid SQL query format"):
        _parse_compound("SELECT id FROM t1 UNION SELECT id ORDER BY id", None)

    with pytest.raises(ValueError, match="LIMIT must be an integer"):
        _parse_compound("SELECT id FROM t1 UNION SELECT id FROM t2 LIMIT 1 EXTRA", None)

    with pytest.raises(ValueError, match="Invalid SQL query format"):
        _parse_compound("SELECT id FROM t1 UNION SELECT id FROM t2 )", None)
    with pytest.raises(ValueError, match="Invalid SQL query format"):
        _parse_compound("SELECT id FROM t1 UNION UNION SELECT id FROM t2", None)
    with pytest.raises(ValueError, match="Invalid SQL query format"):
        _parse_compound("SELECT id FROM t1 UNION (SELECT id FROM t2", None)
    with pytest.raises(ValueError, match="Invalid SQL query format"):
        _parse_compound("SELECT id FROM t1 UNION", None)

    with pytest.raises(ValueError, match="Unsupported SQL action"):
        parse_sql("(SELECT id FROM t1)")


def test_executor_sql_rendering_helper_branches() -> None:
    searched_case = {
        "type": "case",
        "mode": "searched",
        "whens": [
            {"condition": "bad", "result": {"type": "literal", "value": 1}},
            {
                "condition": {"column": "a.id", "operator": "=", "value": 1},
                "result": {"type": "literal", "value": 2},
            },
        ],
    }
    simple_case = {
        "type": "case",
        "mode": "simple",
        "value": {"type": "column", "source": "a", "name": "id"},
        "whens": [{"match": {"type": "literal", "value": 1}, "result": {"type": "literal", "value": "x"}}],
        "else": {"type": "literal", "value": "y"},
    }
    assert "CASE" in SharedExecutor._expression_to_sql(searched_case)
    assert "WHEN" in SharedExecutor._expression_to_sql(simple_case)
    assert SharedExecutor._expression_to_sql(123) == "123"
    assert SharedExecutor._where_operand_to_sql({"type": "literal", "value": 3}, is_column=False) == "3"
    assert SharedExecutor._where_to_sql({"conditions": []}) == ""
    assert "AND" in SharedExecutor._where_to_sql(
        {
            "conditions": [
                {"column": "a.id", "operator": "=", "value": 1},
                {"column": "a.id", "operator": "=", "value": 2},
            ],
            "conjunctions": ["AND"],
        }
    )
    assert "(SUBQUERY)" in SharedExecutor._where_to_sql(
        {"column": "a.id", "operator": "IN", "value": {"type": "subquery"}}
    )
    assert "(3)" in SharedExecutor._where_to_sql(
        {"column": "a.id", "operator": "NOT IN", "value": 3}
    )
    assert "AND" in SharedExecutor._where_to_sql(
        {"column": "a.id", "operator": "BETWEEN", "value": [1, 5]}
    )
    assert SharedExecutor._contains_arithmetic_expression({"type": "alias", "expression": {"type": "alias", "expression": {"type": "literal", "value": 1}}})
    refs = SharedExecutor._collect_expression_column_refs({"type": "alias", "expression": {"type": "column", "source": "a", "name": "id"}})
    assert refs == {"a.id"}


def test_parser_join_validation_and_expression_to_sql_branches() -> None:
    allowed = {"a", "b"}
    _validate_join_column_reference(
        {
            "type": "case",
            "mode": "simple",
            "value": {"type": "column", "source": "a", "name": "id"},
            "whens": [
                "skip",
                {
                    "match": {"type": "column", "source": "a", "name": "id"},
                    "result": {"type": "column", "source": "b", "name": "id"},
                },
            ],
            "else": {"type": "column", "source": "a", "name": "id"},
        },
        allowed,
        "SELECT",
    )
    _validate_join_column_reference(
        {
            "type": "case",
            "mode": "searched",
            "whens": [
                {
                    "condition": {
                        "column": "a.id",
                        "operator": "IN",
                        "value": ({"type": "column", "source": "b", "name": "id"},),
                    },
                    "result": {"type": "column", "source": "a", "name": "id"},
                }
            ],
        },
        allowed,
        "SELECT",
    )
    _validate_join_where_node(
        {
            "column": "a.id",
            "operator": "IN",
            "value": ({"type": "column", "source": "b", "name": "id"},),
        },
        allowed,
    )

    assert _expression_to_sql_for_order_by({"type": "aggregate", "func": "COUNT", "arg": "*"}) == "COUNT(*)"
    assert _expression_to_sql_for_order_by({"type": "literal", "value": "x"}) == "'x'"
    assert _expression_to_sql_for_order_by({"type": "unary_op", "operand": {"type": "literal", "value": 1}}) == "-1"
    assert _expression_to_sql_for_order_by(
        {
            "type": "binary_op",
            "op": "+",
            "left": {"type": "literal", "value": 1},
            "right": {"type": "literal", "value": 2},
        }
    ) == "(1 + 2)"
