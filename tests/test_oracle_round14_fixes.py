from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from excel_dbapi.connection import ExcelConnection
from excel_dbapi.exceptions import ProgrammingError


def _create_round14_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "t"
    sheet.append(["id", "name"])
    sheet.append([1, "a"])
    sheet.append([2, "b"])
    sheet.append([3, "c"])
    workbook.save(path)


def test_case_when_with_window_condition_collects_row_number(tmp_path: Path) -> None:
    file_path = tmp_path / "round14_case_window_condition.xlsx"
    _create_round14_workbook(file_path)

    with ExcelConnection(str(file_path), engine="openpyxl") as conn:
        cursor = conn.cursor()
        cursor.execute(
            "SELECT CASE WHEN ROW_NUMBER() OVER (ORDER BY id) = 1 THEN 'first' ELSE 'other' END AS label FROM t ORDER BY id"
        )
        assert cursor.fetchall() == [("first",), ("other",), ("other",)]


def test_delete_without_placeholders_rejects_extra_parameters(tmp_path: Path) -> None:
    file_path = tmp_path / "round14_delete_extra_params.xlsx"
    _create_round14_workbook(file_path)

    with ExcelConnection(str(file_path), engine="openpyxl") as conn:
        cursor = conn.cursor()
        with pytest.raises(
            ProgrammingError, match="Too many parameters for placeholders"
        ):
            cursor.execute("DELETE FROM t", (123,))


def test_drop_without_placeholders_rejects_extra_parameters(tmp_path: Path) -> None:
    file_path = tmp_path / "round14_drop_extra_params.xlsx"
    _create_round14_workbook(file_path)

    with ExcelConnection(str(file_path), engine="openpyxl") as conn:
        cursor = conn.cursor()
        with pytest.raises(
            ProgrammingError, match="Too many parameters for placeholders"
        ):
            cursor.execute("DROP TABLE t", (123,))


def test_create_without_placeholders_rejects_extra_parameters(tmp_path: Path) -> None:
    file_path = tmp_path / "round14_create_extra_params.xlsx"
    _create_round14_workbook(file_path)

    with ExcelConnection(str(file_path), engine="openpyxl") as conn:
        cursor = conn.cursor()
        with pytest.raises(
            ProgrammingError, match="Too many parameters for placeholders"
        ):
            cursor.execute("CREATE TABLE u (id INTEGER)", (1,))
